"""
A3 pitch poster: reports/MonsoonWatch_Poster_A3.pptx

One portrait A3 slide in plain language: introduction, background,
objectives, a methodology flowchart, four result charts, discussion, future
work and conclusion. Numbers are read from the pipeline's output files (the
weather-free ensemble the web app uses) and the charts are regenerated into
reports/figures/poster/, so re-running keeps the poster in sync.

Run from the project root:  python src/reporting/make_poster.py
"""

import os
import sys
from datetime import timedelta

import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Mm, Pt

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                "..", ".."))
from src.common.epiweek import epiweek_to_end_date  # noqa: E402
from src.common.viz_style import C, apply_style  # noqa: E402

OUT = "reports/MonsoonWatch_Poster_A3.pptx"
FIG = "reports/figures/poster"
FONT = "Arial"

NAVY = "#153e75"
INK = "#1a1a1a"
GREY = "#55545a"
CARD = "#f1f4f8"
WHITE = "#ffffff"
DATA_FILL = "#dce9f8"
MODEL_FILL = "#fde6db"
SDG = [("SDG 3 · Good health", "#4c9f38"), ("SDG 6 · Clean water", "#26bde2"),
       ("SDG 13 · Climate action", "#3f7e44")]

# Page geometry in mm: A3 portrait, outer margin, gap between blocks, card
# padding.
PW, PH = 297, 420
M, G, P = 12, 5, 5
W = PW - 2 * M
MM = 1 / 25.4


# -------------------------------------------------------------------- data
def week_end(year, week):
    return pd.Timestamp(epiweek_to_end_date(int(year), int(week)))


def runs_of(t):
    """Index groups of consecutive weeks, so lines break at missing weeks."""
    t = np.asarray(t)
    return np.split(np.arange(len(t)), np.where(np.diff(t) != 1)[0] + 1)


def load():
    base = pd.read_csv("data/processed/features_base.csv")
    base["t"] = (base.year - 2021) * 52 + base.week - 1
    base["date"] = [week_end(y, w) for y, w in zip(base.year, base.week)]
    # Scores of the weather-free ensemble, from its hold-out backtests (same
    # numbers build_noweather_ensemble.py stores in the model files).
    cmp = pd.read_csv("reports/metrics/model_comparison.csv")
    models = {}
    for h in (1, 2, 3, 4):
        b = pd.read_csv(f"reports/metrics/backtest_noweather_h{h}.csv")
        b = b[np.isfinite(b.pred)]
        err = (b.target - b.pred).abs()
        best = cmp[(cmp.horizon == h) & (cmp.framing == "-")].test_MAE.min()
        models[h] = {"test_MAE": err.mean(),
                     "test_MAPE": (err / b.target).mean() * 100,
                     "best_baseline_MAE": best,
                     "improvement_pct": 100 * (best - err.mean()) / best}
    bt = pd.read_csv("reports/metrics/backtest_noweather_h2.csv")
    bt["t"] = (bt.target_year - 2021) * 52 + bt.target_week - 1
    bt["date"] = [week_end(y, w) for y, w in
                  zip(bt.target_year, bt.target_week)]
    rain = pd.read_csv("reports/metrics/national_rain_test.csv")
    ym = base.groupby("year").age_this_week.mean()
    peak = bt.loc[bt.target.idxmax()]
    return {
        "base": base, "models": models, "bt": bt, "rain": rain,
        "n_weeks": int(base.age_this_week.notna().sum()),
        "mape": [m["test_MAPE"] for m in models.values()],
        "gain": [m["improvement_pct"] for m in models.values()],
        "growth": ym.iloc[-1] / ym.iloc[0],
        "peak": (int(peak.target), int(round(peak.pred, -2))),
    }


# ------------------------------------------------------------------ charts
def new_fig(w_mm, h_mm):
    fig, ax = plt.subplots(figsize=(w_mm * MM, h_mm * MM), layout="constrained")
    return fig, ax


def save(fig, name):
    path = os.path.join(FIG, name)
    fig.savefig(path, dpi=300, bbox_inches=None)
    plt.close(fig)
    return path


def chart_cases(d, w, h):
    b = d["base"].dropna(subset=["age_this_week"])
    fig, ax = new_fig(w, h)
    for y in sorted(d["base"].year.unique()):
        ax.axvspan(week_end(y, 22) - timedelta(days=6), week_end(y, 39),
                   color=C["aqua"], alpha=0.14, lw=0)
    for r in runs_of(b.t):
        ax.plot(b.date.values[r], b.age_this_week.values[r], color=C["blue"],
                lw=1.6)
    top = b.age_this_week.max() * 1.15
    ax.set_ylim(0, top)
    ax.text(week_end(2021, 30), top * 0.97, "monsoon", color="#138a60",
            fontsize=9, ha="center", va="top")
    ax.set_ylabel("Cases per week")
    ax.xaxis.set_major_locator(mdates.YearLocator())
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y"))
    return save(fig, "p1_cases.png")


def chart_forecast(d, w, h):
    bt = d["bt"]
    fig, ax = new_fig(w, h)
    for i, r in enumerate(runs_of(bt.t)):
        ax.plot(bt.date.values[r], bt.target.values[r], color=C["blue"],
                lw=1.8, label="Real cases" if i == 0 else None)
        ax.plot(bt.date.values[r], bt.pred.values[r], color=C["orange"],
                lw=1.8, ls=(0, (4, 2)),
                label="MonsoonWatch forecast" if i == 0 else None)
    ax.set_ylim(0, bt.target.max() * 1.22)
    ax.set_ylabel("Cases per week")
    ax.xaxis.set_major_locator(mdates.MonthLocator(bymonth=[1, 7]))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%b %Y"))
    ax.legend(loc="upper left", ncol=2)
    return save(fig, "p2_forecast.png")


def chart_rules(d, w, h):
    ms = d["models"]
    hs = np.array(sorted(ms))
    rule = [ms[k]["best_baseline_MAE"] for k in hs]
    ours = [ms[k]["test_MAE"] for k in hs]
    fig, ax = new_fig(w, h)
    bw = 0.36
    b1 = ax.bar(hs - bw / 2, rule, bw, color=C["violet"],
                label="Best simple rule")
    b2 = ax.bar(hs + bw / 2, ours, bw, color=C["orange"],
                label="MonsoonWatch")
    for bars in (b1, b2):
        ax.bar_label(bars, fmt="%.0f", padding=2, fontsize=9, color=GREY)
    ax.set_xticks(hs, [f"{k} week{'s' if k > 1 else ''} ahead" for k in hs])
    ax.set_ylim(0, max(rule) * 1.5)
    ax.set_ylabel("Average miss\n(cases per week)")
    ax.grid(axis="x", visible=False)
    ax.legend(loc="upper left", ncol=2)
    return save(fig, "p3_vs_rules.png")


def chart_rain(d, w, h):
    r = d["rain"]
    opts = [("No rainfall (final model)", "no weather", C["orange"]),
            ("Kathmandu Valley rain gauges", "valley gauges (deployed)",
             C["aqua"]),
            ("All-Nepal rainfall (NASA)", "national (NASA)", C["aqua"]),
            ("All-Nepal, wetter/drier than normal",
             "national anomaly (NASA)", C["aqua"])]
    vals = [r[c].mean() for _, c, _ in opts]
    fig, ax = new_fig(w, h)
    bars = ax.barh([o[0] for o in opts], vals, color=[o[2] for o in opts],
                   height=0.62)
    ax.bar_label(bars, fmt="%.1f", padding=3, fontsize=9, color=GREY)
    ax.invert_yaxis()
    ax.set_xlim(0, max(vals) * 1.18)
    ax.set_xlabel("Average miss (cases per week, 1–4 weeks ahead)")
    ax.grid(axis="y", visible=False)
    return save(fig, "p4_rainfall.png")


# -------------------------------------------------------------- pptx bits
def rgb(h):
    return RGBColor.from_string(h.lstrip("#"))


def rect(s, x, y, w, h, fill, radius=None):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = s.shapes.add_shape(shape, Mm(x), Mm(y), Mm(w), Mm(h))
    if radius:
        shp.adjustments[0] = radius / min(w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def fill_text(tf, paras, anchor=MSO_ANCHOR.TOP, margin=0.0):
    """paras: dicts with text or runs=[(text, fmt)], plus size, bold, color,
    align, bullet / number, space_after."""
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Mm(margin)
    tf.margin_top = tf.margin_bottom = Mm(margin)
    for i, sp in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        for txt, fmt in sp.get("runs") or [(sp["text"], {})]:
            r = p.add_run()
            r.text = txt
            f = r.font
            f.name = FONT
            f.size = Pt(fmt.get("size", sp.get("size", 13)))
            f.bold = fmt.get("bold", sp.get("bold", False))
            f.color.rgb = rgb(fmt.get("color", sp.get("color", INK)))
        p.alignment = sp.get("align", PP_ALIGN.LEFT)
        p.space_after = Pt(sp.get("space_after", 4))
        p.line_spacing = 1.0
        if sp.get("bullet") or sp.get("number"):
            pPr = p._p.get_or_add_pPr()
            ind = Mm(5.5)
            pPr.set("marL", str(ind))
            pPr.set("indent", str(-ind))
            clr = pPr.makeelement(qn("a:buClr"), {})
            clr.append(clr.makeelement(qn("a:srgbClr"),
                                       {"val": NAVY.lstrip("#")}))
            pPr.append(clr)
            if sp.get("number"):
                pPr.append(pPr.makeelement(qn("a:buFont"),
                                           {"typeface": FONT}))
                pPr.append(pPr.makeelement(qn("a:buAutoNum"),
                                           {"type": "arabicPeriod"}))
            else:
                pPr.append(pPr.makeelement(qn("a:buFont"),
                                           {"typeface": FONT}))
                pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))


def text(s, x, y, w, h, paras, **kw):
    tb = s.shapes.add_textbox(Mm(x), Mm(y), Mm(w), Mm(h))
    fill_text(tb.text_frame, paras, **kw)
    return tb


def section(s, x, y, w, h, title, paras):
    rect(s, x, y, w, h, CARD, radius=3)
    text(s, x + P, y + 3.5, w - 2 * P, 9,
         [dict(text=title, size=19, bold=True, color=NAVY)])
    if paras:
        text(s, x + P, y + 14, w - 2 * P, h - 14 - P, paras)


def bullets(items, **kw):
    out = []
    for it in items:
        runs = [(it, {})] if isinstance(it, str) else it
        out.append(dict(runs=runs, bullet=True, space_after=5, **kw))
    return out


def B(t):
    return (t, {"bold": True})


def N(t):
    return (t, {})


# -------------------------------------------------------------------- main
def main():
    os.makedirs(FIG, exist_ok=True)
    apply_style()
    plt.rcParams["font.sans-serif"] = ["Arial", "Helvetica", "DejaVu Sans"]
    d = load()
    lo_mape, hi_mape = min(d["mape"]), max(d["mape"])
    lo_gain, hi_gain = min(d["gain"]), max(d["gain"])

    prs = Presentation()
    prs.slide_width, prs.slide_height = Mm(PW), Mm(PH)
    s = prs.slides.add_slide(prs.slide_layouts[6])

    # ---- header
    HDR = 44
    rect(s, 0, 0, PW, HDR, NAVY)
    text(s, M, 3, 210, 22, [dict(text="MonsoonWatch", size=58, bold=True,
                                 color=WHITE)])
    text(s, M, 25, 215, 10, [dict(
        text="Forecasting diarrhoeal disease in Nepal, 1–4 weeks ahead",
        size=21, color="#dbe7f5")])
    text(s, M, 35, 215, 7, [dict(
        text="Pratyush Amatya  ·  Macquarie University  ·  "
             "Pitch for the Planet", size=13, color="#a9c3e6")])
    for i, (label, col) in enumerate(SDG):
        pill = rect(s, PW - M - 52, 6 + i * 11, 52, 9.5, col, radius=2)
        fill_text(pill.text_frame, [dict(text=label, size=12, bold=True,
                                         color=WHITE, align=PP_ALIGN.CENTER,
                                         space_after=0)],
                  anchor=MSO_ANCHOR.MIDDLE)

    # ---- row 1: introduction, background, objectives
    y, h = HDR + G, 64
    cw = (W - 2 * G) / 3
    section(s, M, y, cw, h, "Introduction", [
        dict(text="Every monsoon, floods make drinking water dirty in Nepal. "
                  "Cases of diarrhoeal disease (acute gastroenteritis, AGE) "
                  "shoot up, and young children suffer most.",
             space_after=8),
        dict(runs=[B("A warning a few weeks early"),
                   N(" would give health teams time to get ready.")])])
    section(s, M + cw + G, y, cw, h, "Background", bullets([
        [B("EWARS"), N(", Nepal's warning system, counts cases every week.")],
        "It reports cases only after they happen. It does not forecast.",
        "The weekly numbers are stuck inside PDF reports.",
        "Rain is often blamed for outbreaks. Does it help forecasting?"]))
    section(s, M + 2 * (cw + G), y, cw, h, "Objectives", [
        dict(text=t, number=True, space_after=5) for t in (
            "Turn 5 years of PDF reports into clean weekly data.",
            "Forecast national AGE cases 1–4 weeks ahead.",
            "Test whether rainfall makes forecasts better.",
            "Build a simple web tool that makes the forecast.")])

    # ---- methodology flowchart
    y, h = y + h + G, 56
    section(s, M, y, W, h, "Methodology", None)
    steps = [("Collect", [f"{d['n_weeks']} weekly EWARS", "PDF reports",
                          "(2021–2025)"]),
             ("Extract", ["A program reads", "the PDFs into", "one table"]),
             ("Inputs", ["Recent cases, same", "week last year,",
                         "time of year"]),
             ("Train", ["180 set-ups of 5", "machine-learning",
                        "model types"]),
             ("Combine", ["Average the", "best 15–30", "models"]),
             ("Test", ["Forecast 2024–25", "weeks the models", "never saw"])]
    aw = 7
    bw = (W - 2 * P - 5 * aw) / 6
    by, bh = y + 14, 27
    for i, (name, desc) in enumerate(steps):
        bx = M + P + i * (bw + aw)
        box = rect(s, bx, by, bw, bh, DATA_FILL if i < 3 else MODEL_FILL,
                   radius=2.5)
        fill_text(box.text_frame, [
            dict(text=f"{i + 1}. {name}", size=13, bold=True, color=NAVY,
                 align=PP_ALIGN.CENTER, space_after=2)]
            + [dict(text=line, size=11, align=PP_ALIGN.CENTER, space_after=0)
               for line in desc],
            anchor=MSO_ANCHOR.MIDDLE, margin=1.5)
        if i < 5:
            arr = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                     Mm(bx + bw + 1), Mm(by + bh / 2 - 3),
                                     Mm(aw - 2), Mm(6))
            arr.fill.solid()
            arr.fill.fore_color.rgb = rgb("#9aa3ad")
            arr.line.fill.background()
            arr.shadow.inherit = False
    text(s, M + P, by + bh + 2.5, W - 2 * P, 11, [dict(runs=[
        B("Fair test: "),
        N("each forecast used only data available at that time, and the "
          "models were retrained every week. "),
        B("Rainfall "),
        N("(23 Kathmandu Valley gauges + NASA satellite for all of Nepal) "
          "was tested as an extra input.")], size=12, color=GREY)])

    # ---- results: 2 x 2 charts
    y, h = y + h + G, 128
    section(s, M, y, W, h, "Results", None)
    colw = (W - 2 * P - G) / 2
    cell_h = (h - 14 - P - 4) / 2
    img_h = cell_h - 7
    cells = [("1 · Cases peak every monsoon (green bands)", chart_cases),
             ("2 · Forecast vs real cases, 2 weeks ahead (2024–25)",
              chart_forecast),
             ("3 · Smaller miss than simple rules at every length",
              chart_rules),
             ("4 · Adding rainfall did not reduce the miss", chart_rain)]
    for i, (title, fn) in enumerate(cells):
        cx = M + P + (i % 2) * (colw + G)
        cy = y + 14 + (i // 2) * (cell_h + 4)
        text(s, cx, cy, colw, 7, [dict(text=title, size=13, bold=True)])
        rect(s, cx, cy + 7, colw, img_h, WHITE, radius=2)
        path = fn(d, colw - 2, img_h - 2)
        s.shapes.add_picture(path, Mm(cx + 1), Mm(cy + 8), Mm(colw - 2),
                             Mm(img_h - 2))

    # ---- discussion and future work
    y, h = y + h + G, 64
    hw = (W - G) / 2
    section(s, M, y, hw, h, "Discussion", bullets([
        [B("Past cases and the season"), N(" drive the forecast.")],
        [B("Rainfall did not help"),
         N(", even for all of Nepal. The time of year already covers the "
           "monsoon.")],
        [B("The model type mattered little."),
         N(" Averaging many models and weekly retraining mattered more.")],
        [B("Limits: "),
         N(f"more clinics reporting made counts rise {d['growth']:.0f}×; big "
           f"peaks are under-forecast ({d['peak'][0]} vs ~{d['peak'][1]}).")]
    ]))
    section(s, M + hw + G, y, hw, h, "Future work", bullets([
        [B("Forecast by district"), N(" to show where help is needed most.")],
        [B("Adjust for new health centres"), N(" joining EWARS reporting.")],
        [B("Predict big peaks better"), N(" and add clear alert levels.")],
        [B("Update automatically"), N(" when new EWARS reports come out.")],
        [B("Re-test rainfall"), N(" using district-level cases.")]]))

    # ---- conclusion band with the three headline numbers
    y, h = y + h + G, 30
    rect(s, M, y, W, h, NAVY, radius=3)
    text(s, M + P, y + 3, 150, h - 5, [
        dict(text="Conclusion", size=17, bold=True, color="#ffcf70",
             space_after=2),
        dict(text="MonsoonWatch warns of diarrhoeal disease in Nepal 1–4 "
                  "weeks ahead, using only data Nepal already collects. It "
                  "beats simple rules every time, and needs no weather data.",
             size=12.5, color=WHITE, space_after=0)])
    kpis = [(f"{lo_mape:.0f}–{hi_mape:.0f}%", ["typical", "forecast error"],
             "#ffb38a"),
            (f"{lo_gain:.0f}–{hi_gain:.0f}%", ["better than", "simple rules"],
             "#8fe3c1"),
            ("1–4 wks", ["of early", "warning"], "#9cc7ff")]
    kx0, kw = M + P + 155, (W - 2 * P - 155) / 3
    for i, (num, lab, col) in enumerate(kpis):
        text(s, kx0 + i * kw, y + 3, kw, h - 6, [
            dict(text=num, size=24, bold=True, color=col, space_after=0,
                 align=PP_ALIGN.CENTER),
            *[dict(text=line, size=10.5, color=WHITE, space_after=0,
                   align=PP_ALIGN.CENTER) for line in lab]],
            anchor=MSO_ANCHOR.MIDDLE)

    # ---- footer
    text(s, M, y + h + 1.5, W, 6, [dict(
        text="Data: EDCD EWARS weekly bulletins 2021–2025  ·  DHM rain "
             "gauges  ·  NASA POWER satellite.   Interactive forecast tool: "
             "MonsoonWatch web app.", size=9.5, color=GREY,
        align=PP_ALIGN.CENTER, space_after=0)])

    prs.save(OUT)
    for hz, m in d["models"].items():
        print(f"  h={hz}: MAE {m['test_MAE']:.1f}  MAPE {m['test_MAPE']:.1f}%  "
              f"gain {m['improvement_pct']:+.1f}%")
    print(f"[save] {OUT}  (bottom of content at {y + h + 7.5:.0f} mm of {PH})")


if __name__ == "__main__":
    main()
